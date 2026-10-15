import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_guide_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    doc = Document()
    
    # Page setup - 0.8 inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base Colors
    RED = RGBColor(220, 38, 38)     # #DC2626
    DARK_RED = RGBColor(153, 27, 27) # #991B1B
    GRAY = RGBColor(75, 85, 99)      # #4B5563
    DARK = RGBColor(17, 24, 39)      # #111827
    CODE_COLOR = RGBColor(30, 41, 59)# #1E293B

    lines = content.split("\n")
    i = 0
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i]
        
        # Code block handling
        if line.strip().startswith("```"):
            if in_code_block:
                # End of code block - render single cell table
                in_code_block = False
                code_text = "\n".join(code_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F8F9FA")
                set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(code_text)
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run.font.color.rgb = CODE_COLOR
                # Add spacing after table
                sp_p = doc.add_paragraph()
                sp_p.paragraph_format.space_before = Pt(2)
                sp_p.paragraph_format.space_after = Pt(4)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Markdown Table handling
        if line.strip().startswith("|") and "|" in line.strip()[1:]:
            table_lines = [line]
            while i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
                i += 1
                table_lines.append(lines[i])
            
            # Parse table
            # Filter out separator row (|---|---|)
            parsed_rows = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip().strip("|").split("|")]
                if all(re.match(r"^:?-+:?$", c) for c in cells):
                    continue
                parsed_rows.append(cells)
            
            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                # Normalize row cell count
                for r in parsed_rows:
                    while len(r) < num_cols:
                        r.append("")
                
                table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                for r_idx, row_data in enumerate(parsed_rows):
                    is_header = (r_idx == 0)
                    for c_idx, cell_value in enumerate(row_data):
                        cell = table.cell(r_idx, c_idx)
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        
                        # Strip formatting chars
                        cleaned_val = re.sub(r"\*\*(.*?)\*\*", r"\1", cell_value)
                        cleaned_val = re.sub(r"`(.*?)`", r"\1", cleaned_val)
                        
                        run = p.add_run(cleaned_val)
                        run.font.name = "Calibri"
                        run.font.size = Pt(9)
                        
                        if is_header:
                            set_cell_background(cell, "DC2626")
                            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            bg = "FAF7F2" if r_idx % 2 == 1 else "FFFFFF"
                            set_cell_background(cell, bg)
                            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
                            run.font.color.rgb = DARK
                            
                sp_p = doc.add_paragraph()
                sp_p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # Horizontal Rule
        if line.strip() in ["---", "***", "___"]:
            # Small divider
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("―" * 40)
            run.font.color.rgb = RGBColor(220, 220, 220)
            i += 1
            continue

        # Headings
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[2:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = DARK_RED
            i += 1
            continue

        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(line[3:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = RED
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[4:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = DARK
            i += 1
            continue

        if line.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[5:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(10.5)
            run.bold = True
            run.font.color.rgb = GRAY
            i += 1
            continue

        # Blockquotes / Callout notes
        if line.startswith("> "):
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "FEF2F2") # Light red tint
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[2:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.italic = True
            run.font.color.rgb = DARK_RED
            i += 1
            continue

        # Bullet lists
        if line.strip().startswith("- ") or line.strip().startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            text = line.strip()[2:]
            
            # Parse bold/inline code in bullet
            parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    r.font.color.rgb = DARK
                elif part.startswith("`") and part.endswith("`"):
                    r = p.add_run(part[1:-1])
                    r.font.name = "Consolas"
                    r.font.size = Pt(9)
                    r.font.color.rgb = CODE_COLOR
                else:
                    r = p.add_run(part)
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    r.font.color.rgb = DARK
            i += 1
            continue

        # Numbered lists (e.g. 1. , 2. )
        num_match = re.match(r"^(\d+\.)\s+(.*)", line.strip())
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            text = num_match.group(2)
            parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    r.font.color.rgb = DARK
                elif part.startswith("`") and part.endswith("`"):
                    r = p.add_run(part[1:-1])
                    r.font.name = "Consolas"
                    r.font.size = Pt(9)
                    r.font.color.rgb = CODE_COLOR
                else:
                    r = p.add_run(part)
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    r.font.color.rgb = DARK
            i += 1
            continue

        # Normal text paragraph
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            parts = re.split(r"(\*\*.*?\*\*|`.*?`|\*.*?\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    r.font.color.rgb = DARK
                elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                    r = p.add_run(part[1:-1])
                    r.italic = True
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    r.font.color.rgb = GRAY
                elif part.startswith("`") and part.endswith("`"):
                    r = p.add_run(part[1:-1])
                    r.font.name = "Consolas"
                    r.font.size = Pt(9)
                    r.font.color.rgb = CODE_COLOR
                else:
                    r = p.add_run(part)
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    r.font.color.rgb = DARK

        i += 1

    doc.save(docx_path)
    print(f"Successfully generated Word document at: {docx_path}")

if __name__ == "__main__":
    import os
    src_md = r"d:\New folder (4)\RoadSense-Sensor-Dashboard\ROADSENSE_COMPLETE_GUIDE.md"
    dst_docx = r"d:\New folder (4)\RoadSense-Sensor-Dashboard\ROADSENSE_COMPLETE_GUIDE.docx"
    create_guide_docx(src_md, dst_docx)
    
    # Also copy to nested folder if exists
    nested_dir = r"d:\New folder (4)\RoadSense-Sensor-Dashboard\RoadSense-Sensor-Dashboard"
    if os.path.exists(nested_dir):
        import shutil
        shutil.copy2(dst_docx, os.path.join(nested_dir, "ROADSENSE_COMPLETE_GUIDE.docx"))
        shutil.copy2(src_md, os.path.join(nested_dir, "ROADSENSE_COMPLETE_GUIDE.md"))
        print("Synchronized to nested directory as well.")
